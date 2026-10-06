"""Report generation unit tests (no AI calls; stubbed completions)."""
from __future__ import annotations

import pytest

from owl.reports import generate as gen
from owl.reports.docx_builder import build_docx

TOPICS = [
    {"name": "Attendance", "include": True, "focus": "high"},
    {"name": "Leave", "include": True, "focus": "low"},
    {"name": "Skipped", "include": False, "focus": "high"},
]


def test_word_budgets_follow_focus():
    budgets = gen.word_budgets([t for t in TOPICS if t["include"]], 1000, 2000)
    assert budgets["Attendance"] > budgets["Leave"]
    # Mandated structure overhead: midpoint x1.4 across weights.
    assert sum(budgets.values()) == pytest.approx(2100, abs=150)


def test_split_charts_extracts_and_validates():
    reply = (
        "## Attendance\n\nRate was **83.3% (50 of 60 employees)**.\n\n"
        "```charts\n"
        '[{"type": "bar", "title": "Rate", "labels": ["A", "B"], "values": [83.3, 16.7]},'
        ' {"type": "pie", "title": "Bad", "labels": ["A"], "values": ["x"]}]\n'
        "```"
    )
    md, specs, kpis = gen.split_charts(reply)
    assert "```charts" not in md and "83.3%" in md
    assert len(specs) == 1 and specs[0]["title"] == "Rate"
    assert kpis == []


def test_split_charts_extracts_kpis():
    reply = "Body text.\n\n```kpi\nAttendance rate: 83.3% (50 of 60)\nLate: 10\n```"
    md, specs, kpis = gen.split_charts(reply)
    assert "```kpi" not in md and "Body text." in md
    assert kpis == [("Attendance rate", "83.3% (50 of 60)"), ("Late", "10")]


def test_dedupe_specs_drops_repeats():
    a = {"type": "bar", "title": "One", "labels": ["x"], "values": [1.0]}
    b = {"type": "bar", "title": "Same data", "labels": ["X"], "values": [2.0]}
    c = {"type": "line", "title": "Other", "labels": ["x"], "values": [1.0]}
    assert gen.dedupe_specs([a, b, c], []) == [a, c]
    assert gen.dedupe_specs([b], [a]) == []


def test_mandated_visuals_cover_department_trend_donut():
    from owl.reports.facts import fact_pack
    facts = [fact_pack("att.csv", (
        "department,sex,work_date,status\n"
        "HR,F,2026-06-02,present\nIT,M,2026-06-03,late\n"
        "HR,F,2026-07-04,present\nIT,M,2026-07-05,present\n"
    ))]
    specs = gen.mandated_visuals("Attendance", facts)
    types = {s["type"] for s in specs}
    assert {"bar", "line", "donut"} <= types
    assert specs[0]["title"] == "Attendance by department"


def test_validate_chart_spec_rejects():
    assert gen.validate_chart_spec({"type": "bar"}) is None
    assert gen.validate_chart_spec({
        "type": "bar", "title": "T",
        "labels": ["a", "b"], "values": [1],
    }) is None  # length mismatch
    assert gen.validate_chart_spec({
        "type": "donut", "title": "T",
        "labels": ["a"], "values": [1],
    })["values"] == [1.0]


def test_render_chart_writes_png(tmp_path):
    out = tmp_path / "c.png"
    gen.render_chart({"type": "bar", "title": "T", "labels": ["a"], "values": [1]}, out)
    assert out.exists() and out.stat().st_size > 1000
    out2 = tmp_path / "d.png"
    gen.render_chart({"type": "donut", "title": "T", "labels": ["a", "b"], "values": [70, 30]}, out2)
    assert out2.exists() and out2.stat().st_size > 1000


def test_build_docx_smoke(tmp_path):
    chart = tmp_path / "c.png"
    gen.render_chart({"type": "line", "title": "Trend", "labels": ["W1", "W2"], "values": [80, 90]}, chart)
    md = (
        "## Attendance\n\n"
        "Rate was **83.3% (50 of 60 employees)**.\n\n"
        "- present: 50\n- absent: 10\n\n"
        "| Dept | Rate |\n|---|---|\n| HR | 90% |\n\n"
        "![Trend](build-chart:0)"
    )
    out = tmp_path / "HR_Report.docx"
    build_docx("HR Report", "Sources: a.csv", [("Attendance", md)], [("Trend", chart)], out,
               kpi_boxes=[[("Attendance rate", "83.3% (50 of 60)")]])
    assert out.exists() and out.stat().st_size > 8000
    # TOC field, KPI value, and no duplicated section heading.
    from docx import Document as _Doc
    doc = _Doc(str(out))
    xml = "\n".join(p._p.xml for p in doc.paragraphs)
    assert 'TOC \\o "1-3"' in xml
    assert any("83.3%" in (c.text or "") for t in doc.tables for r in t.rows for c in r.cells)
    h1s = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert h1s.count("Attendance") == 1
